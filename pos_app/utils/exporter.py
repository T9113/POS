import os
from datetime import datetime
from pos_app.config import EXPORTS_DIR
from pos_app.models.settings_model import SettingsModel

class Exporter:
    @staticmethod
    def export_excel(filename: str, title: str, headers: list, rows: list) -> str:
        """Exports tabular data to an Excel (.xlsx) file using openpyxl."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:30]

        # Business info header
        biz_name = SettingsModel.get("business_name", "SwiftPOS")
        ws.append([biz_name])
        ws.append([f"{title} - Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        ws.append([]) # Empty row

        # Style header rows
        title_font = Font(name="Segoe UI", size=16, bold=True, color="2563EB")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        ws["A1"].font = title_font
        ws["A2"].font = sub_font

        # Column headers
        ws.append(headers)
        header_row_idx = 4

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        for row in rows:
            ws.append(row)

        # Style data rows and auto-fit column widths
        row_font = Font(name="Segoe UI", size=10)
        for row_idx in range(header_row_idx + 1, header_row_idx + 1 + len(rows)):
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = row_font
                cell.border = thin_border

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        filepath = os.path.join(EXPORTS_DIR, filename)
        wb.save(filepath)
        return filepath

    @staticmethod
    def export_pdf(filename: str, title: str, headers: list, rows: list, summary_items: list = None) -> str:
        """Exports tabular data to a PDF document using reportlab."""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

        filepath = os.path.join(EXPORTS_DIR, filename)
        doc = SimpleDocTemplate(filepath, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        biz_name = SettingsModel.get("business_name", "SwiftPOS")
        
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=12
        )

        elements.append(Paragraph(f"<b>{biz_name}</b> - {title}", title_style))
        elements.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", subtitle_style))

        # Summary KPIs if provided
        if summary_items:
            kpi_data = [[f"<b>{item[0]}:</b> {item[1]}" for item in summary_items]]
            kpi_table = Table(kpi_data)
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0F172A")),
                ("PADDING", (0, 0), (-1, -1), 6),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]))
            elements.append(kpi_table)
            elements.append(Spacer(1, 14))

        # Main Data Table
        table_data = [headers] + [[str(c) if c is not None else "" for c in row] for row in rows]
        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t)

        doc.build(elements)
        return filepath
